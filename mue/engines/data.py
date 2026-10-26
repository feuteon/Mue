"""Données tabulaires : CSV <-> XLSX <-> JSON (sans Office, via openpyxl)."""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
from pathlib import Path

from . import ConversionError, Task
from ..outputs import OutputCtx
from ..settings import Settings

_INT = re.compile(r"^-?(0|[1-9]\d{0,14})$")
_FLOAT = re.compile(r"^-?\d{1,15}[.,]\d+$")


# --- Lecture -----------------------------------------------------------------------------

def _read_text(src: Path) -> str:
    raw = src.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def read_csv(src: Path) -> tuple[list[list[str]], str]:
    text = _read_text(src)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        if sample.count(";") > sample.count(","):
            dialect = type("D", (csv.excel,), {"delimiter": ";"})
    return [row for row in csv.reader(io.StringIO(text), dialect)], dialect.delimiter


def _typed(v: str, decimal_comma: bool):
    s = v.strip()
    if _INT.match(s):
        return int(s)
    if _FLOAT.match(s):
        if "," in s and not decimal_comma and "." not in s:
            return v
        try:
            return float(s.replace(",", "."))
        except ValueError:
            return v
    return v


def read_xlsx(src: Path) -> dict[str, list[list]]:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    except Exception:
        raise ConversionError("Classeur illisible (seuls .xlsx/.xlsm sont lus sans Office).")
    sheets: dict[str, list[list]] = {}
    active = wb.active.title if wb.active is not None else None
    order = ([active] if active else []) + [n for n in wb.sheetnames if n != active]
    for name in order:
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
        while rows and all(c is None for c in rows[-1]):
            rows.pop()
        if rows:
            width = max(len(r) for r in rows)
            while width and all((len(r) < width or r[width - 1] is None) for r in rows):
                width -= 1
            rows = [r[:width] for r in rows]
        sheets[name] = rows
    wb.close()
    return sheets


def _flatten(obj, prefix="") -> dict:
    out = {}
    if isinstance(obj, dict):
        for k, v in obj.items():
            key = f"{prefix}.{k}" if prefix else str(k)
            if isinstance(v, dict):
                out.update(_flatten(v, key))
            elif isinstance(v, list):
                out[key] = json.dumps(v, ensure_ascii=False)
            else:
                out[key] = v
    else:
        out[prefix or "valeur"] = obj
    return out


def read_json_table(src: Path) -> list[list]:
    try:
        data = json.loads(_read_text(src))
    except ValueError as e:
        raise ConversionError(f"JSON invalide : {e}")
    if isinstance(data, dict):
        lists = [v for v in data.values() if isinstance(v, list)]
        data = lists[0] if len(lists) == 1 else [data]
    if not isinstance(data, list):
        data = [data]
    if data and all(isinstance(r, list) for r in data):
        return [list(r) for r in data]
    flat = [_flatten(r) for r in data]
    keys: list[str] = []
    for r in flat:
        for k in r:
            if k not in keys:
                keys.append(k)
    return [keys] + [[r.get(k) for k in keys] for r in flat]


# --- Écriture ----------------------------------------------------------------------------

def _cell_out(v):
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    return v


def write_csv(rows: list[list], out: Path, delim: str) -> None:
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=delim)
        for r in rows:
            w.writerow(["" if c is None else _cell_out(c) for c in r])


def write_xlsx(sheets: dict[str, list[list]], out: Path) -> None:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "_", name)[:31] or "Feuille")
        for r in rows:
            ws.append([c if not isinstance(c, (dict, list)) else json.dumps(c, ensure_ascii=False) for c in r])
        if rows:
            for c in ws[1]:
                c.font = Font(bold=True)
            ws.freeze_panes = "A2"
            for i in range(1, max(len(r) for r in rows) + 1):
                width = max((len(str(r[i - 1])) for r in rows[:500] if len(r) >= i and r[i - 1] is not None), default=6)
                ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, width + 2))
    wb.save(out)


def rows_to_records(rows: list[list]) -> list[dict]:
    if not rows:
        return []
    header = [str(h) if h not in (None, "") else f"col{i + 1}" for i, h in enumerate(rows[0])]
    return [{header[i]: _cell_out(r[i]) if i < len(r) else None for i in range(len(header))} for r in rows[1:]]


def write_json(obj, out: Path) -> None:
    out.write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def _safe(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*]', "_", name).strip() or "feuille"


# --- Point d'entrée ----------------------------------------------------------------------

def supports(src_ext: str, target: str) -> bool:
    return (src_ext == "csv" and target in ("xlsx", "json")) or \
           (src_ext in ("xlsx", "xlsm") and target in ("csv", "json")) or \
           (src_ext == "json" and target in ("csv", "xlsx"))


def convert(src: Path, target: str, ctx: OutputCtx, task: Task, settings: Settings) -> None:
    ext = src.suffix.lower().lstrip(".")
    delim = settings.csv_delimiter if settings.csv_delimiter in (",", ";", "\t", "|") else ","
    task.progress(0.1)
    if ext == "csv":
        rows, used = read_csv(src)
        dec_comma = used == ";"  # CSV « à la française » : virgule décimale
        if target == "xlsx":
            typed = [rows[0]] + [[_typed(c, dec_comma) for c in r] for r in rows[1:]] if rows else []
            write_xlsx({src.stem[:31]: typed}, ctx.file("xlsx"))
        else:
            write_json(rows_to_records([rows[0]] + [[_typed(c, dec_comma) for c in r] for r in rows[1:]])
                       if rows else [], ctx.file("json"))
    elif ext in ("xlsx", "xlsm"):
        sheets = read_xlsx(src)
        non_empty = {k: v for k, v in sheets.items() if v}
        if not non_empty:
            raise ConversionError("Le classeur est vide.")
        if target == "csv":
            for i, (name, rows) in enumerate(non_empty.items()):
                stem = src.stem if i == 0 else f"{src.stem} - {_safe(name)}"
                write_csv(rows, ctx.file("csv", stem=stem), delim)
        else:
            if len(non_empty) == 1:
                obj = rows_to_records(next(iter(non_empty.values())))
            else:
                obj = {name: rows_to_records(rows) for name, rows in non_empty.items()}
            write_json(obj, ctx.file("json"))
    elif ext == "json":
        rows = read_json_table(src)
        if target == "csv":
            write_csv(rows, ctx.file("csv"), delim)
        else:
            write_xlsx({src.stem[:31]: rows}, ctx.file("xlsx"))
    else:
        raise ConversionError("Format de données non pris en charge.")
    task.progress(1.0)
